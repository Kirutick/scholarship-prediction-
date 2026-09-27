import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')
prs = pptx.Presentation(r'C:\Users\User\Downloads\PBL_Review_1_Scholarship_Eligibility_Prediction (1).pptx')

for i, slide in enumerate(prs.slides, 1):
    print(f"=== SLIDE {i} ===")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    fn = r.font.name
                    fs = r.font.size.pt if r.font.size else None
                    fc = r.font.color.rgb if r.font.color and hasattr(r.font.color, 'rgb') else None
                    print(f"  [{shape.name}] Text: '{r.text[:50]}' | Font: {fn}, Size: {fs}, Color: {fc}")
        elif shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            print(f"  [PICTURE: {shape.name}] (left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height})")
