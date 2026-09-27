"""
Presentation Updater Script for Scholarship Eligibility Prediction.

Updates the existing 10-slide PowerPoint presentation with verified project facts:
1. Slide 1: Title & Team Details (Preserved)
2. Slide 2: Problem Identification & Title Justification (Neutralized scope & wording)
3. Slide 3: Basic Concepts (Clear beginner-friendly definitions)
4. Slide 4: Literature Survey (Separates reference context from synthetic dataset source)
5. Slide 5: Objectives (Objective, neutral wording)
6. Slide 6: Project Planning & Schedule (Preserved W1-W8 timeline & task allocation)
7. Slide 7: Exploratory Data Analysis (EDA: 1,000 records, 63.8%/36.2%, r ≈ -0.57, r ≈ +0.22, updated plots)
8. Slide 8: Model Comparison & Results (93.0% Acc, 0.944 F1, full comparison table, Gini importance chart)
9. Slide 9: References (Distinguishes algorithm/domain references from synthetic dataset)
10. Slide 10: Thank You (Academic ML demonstration label)
"""

import os
import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

sys.stdout.reconfigure(encoding='utf-8')

ORIGINAL_PPT = r"C:\Users\User\Downloads\PBL_Review_1_Scholarship_Eligibility_Prediction_ORIGINAL_BACKUP.pptx"
UPDATED_DOWNLOADS_PPT = r"C:\Users\User\Downloads\PBL_Review_1_Scholarship_Eligibility_Prediction (1).pptx"
UPDATED_WORKSPACE_PPT = r"e:\DS\Scholarship_Eligibility_Prediction_Updated.pptx"
UPDATED_PRESENTATION_DIR_PPT = r"e:\DS\presentation\Scholarship_Eligibility_Prediction.pptx"

COLOR_CYAN = RGBColor(0, 176, 240)    # #00B0F0 from original slides
COLOR_DARK = RGBColor(38, 38, 38)     # #262626 dark gray text
COLOR_MUTED = RGBColor(100, 100, 100) # Gray
COLOR_HIGHLIGHT = RGBColor(230, 245, 255) # Light blue highlight
COLOR_BORDER = RGBColor(180, 210, 240)
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_NAVY = RGBColor(15, 76, 129)    # Deep navy


def format_title(shape, text: str, font_size: float = 34.0):
    """Format slide title with consistent font and color."""
    shape.text = text
    p = shape.text_frame.paragraphs[0]
    p.font.name = "Cambria"
    p.font.size = Pt(font_size)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN


def remove_pictures_from_slide(slide):
    """Remove existing picture shapes from a slide."""
    pic_elements = []
    for shape in slide.shapes:
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            pic_elements.append(shape._element)
    for sp in pic_elements:
        sp.getparent().remove(sp)


def update_slide_2(slide):
    """Update Slide 2: Problem Identification & Title Justification."""
    format_title(slide.shapes[0], "Problem Identification and Title Justification", font_size=32.0)
    tf = slide.shapes[1].text_frame
    tf.clear()

    content = [
        ("Relevance of the Problem", True, 17.0, COLOR_CYAN),
        ("Scholarship screening in higher education traditionally requires checking multiple academic, financial, and demographic criteria across student applications. An automated prediction system can assist institutional review committees in screening applications efficiently based on objective criteria.", False, 13.5, COLOR_DARK),
        ("", False, 6.0, COLOR_DARK),
        ("Problem Statement", True, 17.0, COLOR_CYAN),
        ("Determining scholarship eligibility involves evaluating multiple criteria. This project develops a machine-learning classification pipeline that predicts whether a student application is predicted as Eligible or Not Eligible based on available student attributes.", False, 13.5, COLOR_DARK),
        ("", False, 6.0, COLOR_DARK),
        ("Title Justification and Scope", True, 17.0, COLOR_CYAN),
        ("The title \"Scholarship Eligibility Prediction\" accurately represents the machine-learning classification task using academic and financial attributes. The project scope includes data preprocessing, exploratory data analysis, feature engineering, classification model benchmarking, evaluation, and decision-support prediction, while explicitly excluding official scholarship approval, fund allocation, and document verification.", False, 13.5, COLOR_DARK),
    ]

    for idx, (text, is_bold, size, color) in enumerate(content):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = text
        p.font.name = "Calibri"
        p.font.bold = is_bold
        p.font.size = Pt(size)
        p.font.color.rgb = color


def update_slide_3(slide):
    """Update Slide 3: Basic Concepts."""
    format_title(slide.shapes[0], "Basic Concepts Related to the Project", font_size=34.0)
    tf = slide.shapes[1].text_frame
    tf.clear()

    concepts = [
        ("Machine Learning: ", "Enables computer systems to learn patterns from historical student data to make predictions without being explicitly hardcoded."),
        ("Classification: ", "A supervised learning task that categorizes student applications into discrete classes (Eligible or Not Eligible)."),
        ("Dataset: ", "A structured collection of academic, financial, and demographic student attributes used for model training and evaluation."),
        ("Data Preprocessing: ", "Cleaning, handling categorical encodings, and scaling numerical features into standardized formats for algorithms."),
        ("Feature Selection / Engineering: ", "Identifying and transforming key predictor attributes (e.g., family income, board marks) that contribute to predictions."),
        ("Model Training: ", "Fitting classification algorithms on training data using supervised learning to minimize prediction error."),
        ("Prediction System: ", "Evaluating new student attribute inputs through the trained pipeline to output predicted eligibility with confidence scores.")
    ]

    for idx, (label, desc) in enumerate(concepts):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(8.0)
        
        r1 = p.add_run()
        r1.text = label
        r1.font.name = "Calibri"
        r1.font.bold = True
        r1.font.size = Pt(14.0)
        r1.font.color.rgb = COLOR_CYAN

        r2 = p.add_run()
        r2.text = desc
        r2.font.name = "Calibri"
        r2.font.bold = False
        r2.font.size = Pt(13.0)
        r2.font.color.rgb = COLOR_DARK


def update_slide_4(slide):
    """Update Slide 4: Literature Survey."""
    format_title(slide.shapes[0], "Literature Survey & Research Gap", font_size=34.0)
    tf = slide.shapes[1].text_frame
    tf.clear()

    lines = [
        ("Related Systems Survey:", True, 15.0, COLOR_CYAN),
        ("  1. Scholarship Recommender: Suggests scholarships | Limitation: No eligibility prediction", False, 12.0, COLOR_DARK),
        ("  2. Financial Aid Predictor: Uses income data | Limitation: Limited academic factors considered", False, 12.0, COLOR_DARK),
        ("  3. AI Screening System: Faster evaluation | Limitation: Requires very large enterprise datasets", False, 12.0, COLOR_DARK),
        ("  4. Student Support System: Assists decisions | Limitation: Less personalized to individual criteria", False, 12.0, COLOR_DARK),
        ("  5. ML Eligibility Model: Good initial accuracy | Limitation: Restricted feature scope", False, 12.0, COLOR_DARK),
        ("", False, 4.0, COLOR_DARK),
        ("Research Gap Addressed:", True, 15.0, COLOR_CYAN),
        ("  • Multidimensional Feature Integration: Combines academic, financial, and demographic attributes in a unified pipeline rather than isolated evaluations.", False, 12.0, COLOR_DARK),
        ("  • Multi-Algorithm Benchmarking: Systematically compares four distinct classification paradigms (linear, tree, ensemble, probabilistic).", False, 12.0, COLOR_DARK),
        ("  • Zero-Leakage Pipeline: Implements strict ColumnTransformer preprocessing fitted exclusively on training data.", False, 12.0, COLOR_DARK),
        ("", False, 4.0, COLOR_DARK),
        ("Reference Context Note:", True, 13.0, COLOR_MUTED),
        ("  • Government scholarship guidelines and AISHE statistics provide conceptual domain context; the dataset is a synthetic reference benchmark created for academic experimentation.", False, 11.5, COLOR_MUTED)
    ]

    for idx, (text, is_bold, size, color) in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = text
        p.font.name = "Calibri"
        p.font.bold = is_bold
        p.font.size = Pt(size)
        p.font.color.rgb = color


def update_slide_5(slide):
    """Update Slide 5: Objectives."""
    format_title(slide.shapes[0], "Objectives of the Project", font_size=34.0)
    tf = slide.shapes[1].text_frame
    tf.clear()

    objectives = [
        "1. Develop a machine learning classification model for scholarship eligibility prediction.",
        "2. Analyze academic, financial, and demographic student attributes affecting eligibility.",
        "3. Perform structured data preprocessing and exploratory data analysis (EDA).",
        "4. Train and benchmark multiple classification algorithms (Logistic Regression, Decision Tree, Random Forest, Naïve Bayes).",
        "5. Evaluate model performance using standard evaluation metrics (Accuracy, Precision, Recall, F1-Score) to predict students labeled as eligible in the dataset.",
        "6. Provide a decision-support prediction interface with probability estimation."
    ]

    for idx, obj in enumerate(objectives):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = obj
        p.space_after = Pt(14.0)
        p.font.name = "Calibri"
        p.font.bold = (idx == 0 or idx == 3)
        p.font.size = Pt(15.0)
        p.font.color.rgb = COLOR_DARK


def update_slide_7(slide):
    """Update Slide 7: Exploratory Data Analysis (EDA)."""
    format_title(slide.shapes[0], "Exploratory Data Analysis (EDA)", font_size=32.0)
    tf = slide.shapes[1].text_frame
    tf.clear()

    eda_points = [
        ("Dataset: ", "1,000 synthetic student records (Community, FamilyIncome, 12thMarks, FirstGraduate, District, CollegeType, Course)."),
        ("Target Distribution: ", "63.8% Eligible (638) vs 36.2% Not Eligible (362) — moderately imbalanced classification target."),
        ("FamilyIncome Correlation: ", "r ≈ -0.57 (Calculated r = -0.5739) — shows the strongest negative linear association with eligibility."),
        ("12thMarks Correlation: ", "r ≈ +0.22 (Calculated r = +0.2226) — shows a moderate positive linear association with eligibility."),
        ("Causality Note: ", "Correlation indicates linear statistical association within the dataset, not causation (income/marks do not independently cause eligibility).")
    ]

    for idx, (label, text) in enumerate(eda_points):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(4.0)
        
        r1 = p.add_run()
        r1.text = label
        r1.font.name = "Calibri"
        r1.font.bold = True
        r1.font.size = Pt(12.5)
        r1.font.color.rgb = COLOR_CYAN if "Correlation" in label else COLOR_DARK

        r2 = p.add_run()
        r2.text = text
        r2.font.name = "Calibri"
        r2.font.bold = False
        r2.font.size = Pt(12.0)
        r2.font.color.rgb = COLOR_MUTED if "Causality" in label else COLOR_DARK

    # Replace pictures with verified fresh plots
    remove_pictures_from_slide(slide)

    # Left: Community Distribution Chart
    pic1_path = "outputs/plots/community_distribution.png"
    # Right: Income Distribution Chart
    pic2_path = "outputs/plots/family_income_distribution.png"

    if os.path.exists(pic1_path):
        slide.shapes.add_picture(pic1_path, Inches(0.6), Inches(3.45), Inches(5.8), Inches(3.6))
    if os.path.exists(pic2_path):
        slide.shapes.add_picture(pic2_path, Inches(6.8), Inches(3.45), Inches(5.9), Inches(3.6))


def update_slide_8(slide):
    """Update Slide 8: Model Comparison & Results."""
    format_title(slide.shapes[0], "Model Comparison & Results", font_size=32.0)
    tf = slide.shapes[1].text_frame
    tf.clear()

    summary_points = [
        "Four classification algorithms were trained on 800 samples and evaluated on 200 held-out test samples (stratified 80/20 split).",
        "Random Forest achieved 93.0% accuracy and 0.944 F1-score on the held-out test set, performing best in the experiment.",
        "Random Forest Feature Importance: FamilyIncome (46.94%) and 12thMarks (21.28%) are the top predictors in the model."
    ]

    for idx, text in enumerate(summary_points):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = "• " + text
        p.space_after = Pt(3.0)
        p.font.name = "Calibri"
        p.font.bold = (idx == 1)
        p.font.size = Pt(12.5)
        p.font.color.rgb = COLOR_CYAN if idx == 1 else COLOR_DARK

    remove_pictures_from_slide(slide)

    # 1. Native PowerPoint Table on the Left
    table_left = Inches(0.5)
    table_top = Inches(2.7)
    table_width = Inches(6.4)
    table_height = Inches(4.3)

    shape_tbl = slide.shapes.add_table(5, 5, table_left, table_top, table_width, table_height)
    tbl = shape_tbl.table

    tbl.columns[0].width = Inches(2.4)
    tbl.columns[1].width = Inches(1.0)
    tbl.columns[2].width = Inches(1.0)
    tbl.columns[3].width = Inches(1.0)
    tbl.columns[4].width = Inches(1.0)

    headers = ["Model", "Accuracy", "Precision", "Recall", "F1-Score"]
    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_NAVY
        p = cell.text_frame.paragraphs[0]
        p.font.name = "Calibri"
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_WHITE
        p.alignment = PP_ALIGN.CENTER if j > 0 else PP_ALIGN.LEFT

    rows_data = [
        ("Logistic Regression", "86.0%", "0.8906", "0.8906", "0.8906", False),
        ("Decision Tree", "90.5%", "0.9504", "0.8984", "0.9237", False),
        ("Random Forest (Best)", "93.0%", "0.9672", "0.9219", "0.9440", True),
        ("Naïve Bayes", "84.0%", "0.8810", "0.8672", "0.8740", False)
    ]

    for i, (m_name, acc, prec, rec, f1, is_highlight) in enumerate(rows_data, 1):
        row_vals = [m_name, acc, prec, rec, f1]
        for j, val in enumerate(row_vals):
            cell = tbl.cell(i, j)
            cell.text = val
            cell.fill.solid()
            if is_highlight:
                cell.fill.fore_color.rgb = COLOR_HIGHLIGHT
            else:
                cell.fill.fore_color.rgb = COLOR_WHITE if i % 2 == 1 else RGBColor(245, 248, 252)

            p = cell.text_frame.paragraphs[0]
            p.font.name = "Calibri"
            p.font.bold = is_highlight
            p.font.size = Pt(11.5)
            p.font.color.rgb = COLOR_NAVY if is_highlight else COLOR_DARK
            p.alignment = PP_ALIGN.CENTER if j > 0 else PP_ALIGN.LEFT

    # 2. Right Picture: Feature Importance Chart
    pic_fi_path = "outputs/plots/random_forest_feature_importance.png"
    if os.path.exists(pic_fi_path):
        slide.shapes.add_picture(pic_fi_path, Inches(7.1), Inches(2.65), Inches(5.8), Inches(4.35))


def update_slide_9(slide):
    """Update Slide 9: References."""
    format_title(slide.shapes[0], "References", font_size=36.0)
    tf = slide.shapes[1].text_frame
    tf.clear()

    refs = [
        ("Pedregosa, F., et al. (2011). ", "Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825-2830."),
        ("Breiman, L. (2001). ", "Random Forests. Machine Learning, 45(1), 5-32."),
        ("Quinlan, J. R. (1986). ", "Induction of Decision Trees. Machine Learning, 1(1), 81-106."),
        ("McKinney, W. (2010). ", "Data Structures for Statistical Computing in Python (pandas). Proceedings of the 9th Python in Science Conference."),
        ("Government of Tamil Nadu: ", "Post-Matric Scholarship Scheme Guidelines (Adi Dravidar & Tribal Welfare / Backward Classes Welfare Department) – Used as conceptual reference context for eligibility criteria."),
        ("Ministry of Education, Government of India: ", "All India Survey on Higher Education (AISHE) Reports – Used as reference context for higher education enrolment demographics.")
    ]

    for idx, (auth, title) in enumerate(refs):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(8.0)
        
        r1 = p.add_run()
        r1.text = auth
        r1.font.name = "Calibri"
        r1.font.bold = True
        r1.font.size = Pt(12.5)
        r1.font.color.rgb = COLOR_DARK

        r2 = p.add_run()
        r2.text = title
        r2.font.name = "Calibri"
        r2.font.bold = False
        r2.font.size = Pt(12.0)
        r2.font.color.rgb = COLOR_DARK

    # Clarification note at bottom
    p_note = tf.add_paragraph()
    p_note.space_before = Pt(8.0)
    r_n = p_note.add_run()
    r_n.text = "Note: References provide algorithmic background and conceptual domain context; the project dataset is a synthetic reference dataset generated for academic machine-learning experimentation."
    r_n.font.name = "Calibri"
    r_n.font.italic = True
    r_n.font.size = Pt(11.0)
    r_n.font.color.rgb = COLOR_MUTED


def update_slide_10(slide):
    """Update Slide 10: Thank You."""
    format_title(slide.shapes[0], "Thank You!", font_size=56.0)
    p = slide.shapes[0].text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER

    # Add subtitle box
    sub_box = slide.shapes.add_textbox(Inches(2.0), Inches(4.3), Inches(9.3), Inches(1.8))
    tf = sub_box.text_frame
    
    p1 = tf.paragraphs[0]
    p1.text = "Academic Machine Learning Classification Demonstration"
    p1.font.name = "Calibri"
    p1.font.bold = True
    p1.font.size = Pt(22.0)
    p1.font.color.rgb = COLOR_DARK
    p1.alignment = PP_ALIGN.CENTER
    p1.space_after = Pt(6.0)

    p2 = tf.add_paragraph()
    p2.text = "Course: AD4V71 – Data Science\nTeam: Kirutick Siddhesh V (210425243120) & Krishna H (210425243124)"
    p2.font.name = "Calibri"
    p2.font.size = Pt(16.0)
    p2.font.color.rgb = COLOR_MUTED
    p2.alignment = PP_ALIGN.CENTER


def main():
    print("=" * 65)
    print("UPDATING POWERPOINT PRESENTATION WITH VERIFIED FACTS")
    print("=" * 65)

    if not os.path.exists(ORIGINAL_PPT):
        print(f"Error: Original PPT backup not found at {ORIGINAL_PPT}")
        return

    prs = pptx.Presentation(ORIGINAL_PPT)
    print(f"Loaded presentation: {len(prs.slides)} slides.")

    # Apply updates slide-by-slide
    print("[UPDATING] Slide 2: Problem Identification & Title Justification...")
    update_slide_2(prs.slides[1])

    print("[UPDATING] Slide 3: Basic Concepts...")
    update_slide_3(prs.slides[2])

    print("[UPDATING] Slide 4: Literature Survey & Context Note...")
    update_slide_4(prs.slides[3])

    print("[UPDATING] Slide 5: Objectives (Objective terminology)...")
    update_slide_5(prs.slides[4])

    print("[VERIFYING] Slide 6: Project Planning & Schedule preserved.")

    print("[UPDATING] Slide 7: Exploratory Data Analysis (EDA)...")
    update_slide_7(prs.slides[6])

    print("[UPDATING] Slide 8: Model Comparison Table & Results (93.0% / 0.944)...")
    update_slide_8(prs.slides[7])

    print("[UPDATING] Slide 9: References (Distinguishing dataset from context)...")
    update_slide_9(prs.slides[8])

    print("[UPDATING] Slide 10: Thank You (Academic ML Demonstration)...")
    update_slide_10(prs.slides[9])

    # Save to all target locations
    os.makedirs(os.path.dirname(UPDATED_PRESENTATION_DIR_PPT), exist_ok=True)

    targets = [
        UPDATED_WORKSPACE_PPT,
        UPDATED_PRESENTATION_DIR_PPT,
        UPDATED_DOWNLOADS_PPT
    ]

    for target in targets:
        prs.save(target)
        print(f"[SAVED] Updated presentation -> {target}")

    print("=" * 65)
    print("PRESENTATION UPDATE COMPLETED SUCCESSFULLY!")
    print("=" * 65)


if __name__ == "__main__":
    main()
